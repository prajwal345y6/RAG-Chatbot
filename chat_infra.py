import os
import pickle
import torch
import numpy as np
import hashlib
import faiss
from flask import Flask, render_template, request
from fuzzywuzzy import fuzz
from transformers import AutoTokenizer, LlamaForCausalLM
from office365.sharepoint.client_context import ClientContext
from office365.runtime.auth.user_credential import UserCredential
from io import BytesIO
from docx import Document
import fitz
from datetime import datetime, timedelta
import base64
from PIL import Image
from io import BytesIO

# Flask App
app = Flask(__name__)

# SharePoint Credentials
SITE_URL = "https://myldev.sharepoint.com/sites/otptest"
USERNAME = "Ramesh@myldev.onmicrosoft.com"
PASSWORD = "Job28124"

# FAISS Index Storage
INDEX_FOLDER = "faiss_indices"
INDEX_DATA_FILE = "index_metadata.pkl"

# Stop Words
STOP_WORDS = set(["the", "and", "is", "to", "for", "of", "in", "on", "with", "a", "an", "procedures", "process", "instructions", "details", "how", "can", "I", "i", "explain", "it", "this", "achieve", "that", "by", "as", "at", "be", "or", "from", "provide", "give", "me", "steps"])


# Load LLaMA model and tokenizer
device = 'cuda' if torch.cuda.is_available() else 'cpu'
model_name = "meta-llama/Llama-2-7b-hf"
tokenizer = AutoTokenizer.from_pretrained(model_name)
model = LlamaForCausalLM.from_pretrained(model_name, device_map="auto",
                                         torch_dtype=torch.float16 if device == "cuda" else torch.float32)
model.to(device)

if tokenizer.pad_token is None:
    tokenizer.add_special_tokens({'pad_token': tokenizer.eos_token})
    model.resize_token_embeddings(len(tokenizer))

# Helper Functions
#def preprocess_text(text):
    #return ''.join(c.lower() if c.isalnum() or c.isspace() else ' ' for c in text).strip()
# Helper Functions
def preprocess_text(text):
    words = text.lower().split()
    filtered_words = [word for word in words if word not in STOP_WORDS]
    return ' '.join(filtered_words).strip()


def encode_text_with_llama(text):
    inputs = tokenizer(text, return_tensors="pt", padding=True, truncation=True, max_length=512).to(device)
    with torch.no_grad():
        outputs = model(**inputs, output_hidden_states=True)
        embeddings = outputs.hidden_states[-1].mean(dim=1).squeeze().cpu().numpy()
    return embeddings.astype(np.float32)


def add_document_to_faiss(file_name, content, embedding, index, index_data):
    doc_id = hashlib.sha256(file_name.encode('utf-8')).hexdigest()
    metadata = {"file_name": file_name, "full_content": content}

    index_data.append((doc_id, metadata, embedding))
    index.add(np.array([embedding], dtype=np.float32))

    print(f"Document '{file_name}' indexed successfully.")


def check_for_document_modifications(ctx, folder_url, indexed_folders, time_threshold_minutes=1):
    folder = ctx.web.get_folder_by_server_relative_url(folder_url)
    files = folder.files
    ctx.load(files)
    ctx.execute_query()

    current_time = datetime.utcnow()

    for file in files:
        file_name = file.properties["Name"]
        last_modified_time = file.properties["TimeLastModified"]

        if isinstance(last_modified_time, datetime):
            time_difference = current_time - last_modified_time

            if time_difference <= timedelta(minutes=time_threshold_minutes):
                print(f"Document '{file_name}' was modified recently. Updating FAISS index...")

                # Download the modified file
                file_content = download_file_from_sharepoint(ctx, file.properties["ServerRelativeUrl"])

                # Extract content
                if file_name.endswith(".docx"):
                    content = read_docx(file_content)
                elif file_name.endswith(".pdf"):
                    content = read_pdf(file_content)
                else:
                    print(f"Unsupported file format: {file_name}")
                    continue

                full_text = ' '.join(content)
                preprocessed_text = preprocess_text(full_text)
                new_embedding = encode_text_with_llama(preprocessed_text)

                # Update FAISS
                if folder_url in indexed_folders:
                    index, index_data = indexed_folders[folder_url]

                    # Remove old embedding (if exists)
                    index_data = [entry for entry in index_data if entry[1]['file_name'] != file_name]
                    index.reset()  # Clear existing index
                    for _, _, embedding in index_data:
                        index.add(np.array([embedding], dtype=np.float32))

                    # Add new embedding
                    add_document_to_faiss(file_name, content, new_embedding, index, index_data)
                    indexed_folders[folder_url] = (index, index_data)

                    # Save updated index
                    save_indexed_data(indexed_folders)
                    print(f"FAISS index updated for '{file_name}'")
                else:
                    print(f"Folder '{folder_url}' not found in indexed data. Skipping reindexing.")


def create_faiss_index(dimension=4096):
    return faiss.IndexFlatL2(dimension)


def read_pdf(file_content):
    doc = fitz.open(stream=file_content, filetype="pdf")
    return [page.get_text("text").strip() for page in doc if page.get_text("text").strip()]


def read_docx(file_content):
    doc = Document(BytesIO(file_content))
    return [para.text.strip() for para in doc.paragraphs if para.text.strip()]


def download_file_from_sharepoint(ctx, file_url):
    file = ctx.web.get_file_by_server_relative_url(file_url)
    ctx.load(file)
    ctx.execute_query()
    file_content_stream = BytesIO()
    file.download(file_content_stream)
    ctx.execute_query()
    file_content_stream.seek(0)
    return file_content_stream.read()


# Function to extract images from PDF
def extract_images_from_pdf(file_content):
    images = []
    doc = fitz.open(stream=file_content, filetype="pdf")

    for page_index in range(len(doc)):
        for img_index, img in enumerate(doc[page_index].get_images(full=True)):
            xref = img[0]
            base_image = doc.extract_image(xref)
            image_bytes = base_image["image"]

            # Convert to base64
            img_base64 = base64.b64encode(image_bytes).decode("utf-8")
            images.append(img_base64)

    print(f"Extracted {len(images)} images from PDF.")  # Debugging statement
    return images


# Function to extract images from DOCX
def extract_images_from_docx(file_content):
    images = []
    doc = Document(BytesIO(file_content))

    for rel_id, rel in doc.part.rels.items():
        try:
            # Skip external links to avoid the ValueError
            if hasattr(rel, "target_mode") and rel.target_mode == "External":
                continue

            # Ensure it's an image before extracting
            if "image" in rel.target_ref:
                image_data = rel.target_part.blob  # Extract image data
                img_base64 = base64.b64encode(image_data).decode("utf-8")
                images.append(img_base64)

        except (AttributeError, KeyError, ValueError):
            # Skip if relationship does not have a valid target_part (external images)
            continue

    print(f"Extracted {len(images)} images from DOCX.")  # Debugging statement
    return images


# Modify index_documents_in_faiss to extract images
def index_documents_in_faiss(ctx, folder_url, index, index_data):
    folder = ctx.web.get_folder_by_server_relative_url(folder_url)
    files = folder.files
    ctx.load(files)
    ctx.execute_query()

    for file in files:
        file_name = file.properties["Name"]
        file_url = file.properties["ServerRelativeUrl"]
        file_content = download_file_from_sharepoint(ctx, file_url)

        images = []  # Store extracted images

        if file_name.endswith(".docx"):
            content = read_docx(file_content)
            images = extract_images_from_docx(file_content)  # Extract images
        elif file_name.endswith(".pdf"):
            content = read_pdf(file_content)
            images = extract_images_from_pdf(file_content)  # Extract images
        else:
            continue

        full_text = ' '.join(content)
        preprocessed_text = preprocess_text(full_text)
        embedding = encode_text_with_llama(preprocessed_text)

        doc_id = hashlib.sha256(file_name.encode('utf-8')).hexdigest()
        metadata = {"file_name": file_name, "full_content": content, "images": images}  # Store images in metadata

        index_data.append((doc_id, metadata, embedding))
        index.add(np.array([embedding], dtype=np.float32))

        print(f"Document '{file_name}' indexed successfully with {len(images)} images.")


def save_indexed_data(indexed_folders):
    os.makedirs(INDEX_FOLDER, exist_ok=True)
    index_metadata = {}
    for folder_name, (index, index_data) in indexed_folders.items():
        faiss.write_index(index, os.path.join(INDEX_FOLDER, f"{folder_name}.faiss"))
        index_metadata[folder_name] = index_data
    with open(INDEX_DATA_FILE, "wb") as f:
        pickle.dump(index_metadata, f)


def load_indexed_data():
    if not os.path.exists(INDEX_DATA_FILE):
        return {}
    with open(INDEX_DATA_FILE, "rb") as f:
        index_metadata = pickle.load(f)
    indexed_folders = {}
    for folder_name, index_data in index_metadata.items():
        index_path = os.path.join(INDEX_FOLDER, f"{folder_name}.faiss")
        if os.path.exists(index_path):
            index = faiss.read_index(index_path)
            indexed_folders[folder_name] = (index, index_data)
    return indexed_folders



# Flask Routes
@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        folder_name = request.form["folder_name"]
        user_query = request.form["query"]

        indexed_folders = load_indexed_data()

        if folder_name not in indexed_folders:
            print(f"Indexing new folder: {folder_name}")
            index_data = []
            index = create_faiss_index(dimension=4096)
            ctx = ClientContext(SITE_URL).with_credentials(UserCredential(USERNAME, PASSWORD))
            index_documents_in_faiss(ctx, f"/sites/otptest/Shared%20Documents/{folder_name}", index, index_data)
            indexed_folders[folder_name] = (index, index_data)
            save_indexed_data(indexed_folders)
        else:
            print(f"Using existing index for folder '{folder_name}'.")
            index, index_data = indexed_folders[folder_name]

        preprocessed_query = preprocess_text(user_query)

        best_match_score = 0
        best_match_line = None
        best_match_document = None
        best_match_index = -1
        best_match_images = []  # Store images for best match

        for doc_id, metadata, _ in index_data:
            content = metadata['full_content']
            for i, line in enumerate(content):
                similarity = fuzz.ratio(preprocessed_query, preprocess_text(line))
                if similarity > best_match_score:
                    best_match_score = similarity
                    best_match_line = line
                    best_match_document = metadata['file_name']
                    best_match_index = i
                    best_match_images = metadata.get("images", [])  # Get images for this document

        print(f"Found {len(best_match_images)} images for query: {user_query}")  # Debugging statement

        SIMILARITY_THRESHOLD = 70  # Adjust this threshold as needed

        # Identify file type
        file_extension = best_match_document.split(".")[-1].lower() if best_match_document else ""

        # Apply similarity threshold **only for .docx files**
        if (file_extension == "docx" and best_match_score >= SIMILARITY_THRESHOLD) or file_extension == "pdf":
            best_doc_content = next(
                md['full_content'] for doc_id, md, _ in index_data if md['file_name'] == best_match_document)
            matched_lines = best_doc_content[best_match_index:best_match_index + 41]

            return render_template("index.html", results="\n".join(matched_lines), images=best_match_images)

        # Case: No high similarity score for DOCX but a valid match exists
        #elif file_extension == "docx" and best_match_index != -1 and best_match_document:
        #    best_doc_content = next(
        #        md['full_content'] for doc_id, md, _ in index_data if md['file_name'] == best_match_document)
        #    matched_lines = best_doc_content[best_match_index:best_match_index + 41]
#
        #    return render_template("index.html",
        #                           results="Best available match:\n" + "\n".join(matched_lines),
        #                           images=best_match_images)

        # Case: No match found
        return render_template("index.html", error="No results found for this query.", images=[])

    return render_template("index.html")


if __name__ == "__main__":
    app.run(debug=True, host='0.0.0.0', port=5000, use_reloader=False)
